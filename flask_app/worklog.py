"""
GENRI LABS — OPT WORK LOG

Contemporaneous record of Genri Labs work for OPT self-employment:
one row per work block, totaled per week (Mon–Sun) against the
20-hour minimum. Same columns as the Excel timesheet, so exports
drop straight into the OPT evidence folder.

Two ways rows arrive:
  - "manual": typed in the admin at /worklog → approved immediately.
  - "claude": posted by Claude to /api/worklog/ingest with the
    WORKLOG_API_KEY → saved as a DRAFT with no hours. Drafts never
    count toward weekly totals until Glad sets the hours and approves.

The table creates itself on first use (same pattern as projects.py).
"""
import datetime
import io
from decimal import Decimal, InvalidOperation

from db import query, execute

# Same dropdown as Genri_Labs_OPT_Timesheet.xlsx
CATEGORIES = (
    "Software Development",
    "Product/Project Management",
    "Business Analysis & Planning",
    "Technical Documentation",
    "Website/Infrastructure",
    "Research (Tech/AI)",
    "Business Development & Outreach",
    "Marketing/Content",
    "Administration & Compliance",
)
STATUSES = ("draft", "approved")
WEEKLY_MINIMUM = 20
FIRST_WEEK = datetime.date(2026, 6, 29)  # Monday of the week Genri Labs LLC was formed
MAX_HOURS_PER_ROW = 16
MAX_ACTIVITY = 4000

_CREATE_SQL = """
CREATE TABLE IF NOT EXISTS work_sessions (
    id          INT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
    work_date   DATE          NOT NULL,
    hours       DECIMAL(5,2)  NULL,                      -- NULL until a draft is approved
    category    VARCHAR(64)   NOT NULL,
    activity    TEXT          NOT NULL,                   -- a day's summary can be long
    project     VARCHAR(120)  NOT NULL DEFAULT '',
    evidence    VARCHAR(500)  NOT NULL DEFAULT '',
    degree_link VARCHAR(500)  NOT NULL DEFAULT '',          -- optional: how the work applies the MBA-MIS
    notes       VARCHAR(500)  NOT NULL DEFAULT '',
    status      VARCHAR(16)   NOT NULL DEFAULT 'approved',  -- 'draft' | 'approved'
    source      VARCHAR(16)   NOT NULL DEFAULT 'manual',    -- 'manual' | 'claude'
    created_at  TIMESTAMP     NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at  TIMESTAMP     NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
    INDEX idx_work_date (work_date)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
"""

_COLUMNS = ("id, work_date, hours, category, activity, project, evidence, degree_link, notes, "
            "status, source, created_at")

_ready = False


def ensure_table():
    """Create the work_sessions table (and add later columns) once per process."""
    global _ready
    if _ready:
        return
    execute(_CREATE_SQL)
    # Tables created before degree_link existed get the column added in place.
    if not query("SHOW COLUMNS FROM work_sessions LIKE %s", ("degree_link",)):
        execute("ALTER TABLE work_sessions ADD COLUMN degree_link VARCHAR(500) "
                "NOT NULL DEFAULT '' AFTER evidence")
    # Daily summaries outgrow the original VARCHAR(1000).
    col = query("SHOW COLUMNS FROM work_sessions LIKE %s", ("activity",), one=True)
    if col and not str(col["Type"]).lower().startswith("text"):
        execute("ALTER TABLE work_sessions MODIFY activity TEXT NOT NULL")
    _ready = True


def week_start(day):
    """Monday of the week containing `day`."""
    return day - datetime.timedelta(days=day.weekday())


def _as_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return datetime.date.fromisoformat(str(value))


def _serialize(row):
    day = _as_date(row["work_date"])
    return {
        "id": row["id"],
        "work_date": day.isoformat(),
        "week_start": week_start(day).isoformat(),
        "hours": float(row["hours"]) if row["hours"] is not None else None,
        "category": row["category"],
        "activity": row["activity"],
        "project": row["project"],
        "evidence": row["evidence"],
        "degree_link": row["degree_link"],
        "notes": row["notes"],
        "status": row["status"],
        "source": row["source"],
    }


def list_sessions():
    ensure_table()
    rows = query(f"SELECT {_COLUMNS} FROM work_sessions ORDER BY work_date DESC, id DESC")
    return [_serialize(r) for r in rows]


def clean(data, require_hours=True):
    """Validate a work-session payload. Returns (fields, error_message)."""
    try:
        work_date = datetime.date.fromisoformat((data.get("work_date") or "").strip())
    except ValueError:
        return None, "Date is required (YYYY-MM-DD)."
    if work_date > datetime.date.today() + datetime.timedelta(days=1):
        return None, "Date can't be in the future."

    category = (data.get("category") or "").strip()
    if category not in CATEGORIES:
        return None, "Pick a category from the list."

    activity = (data.get("activity") or "").strip()
    if not activity:
        return None, "Describe what you did."

    raw_hours = data.get("hours")
    hours = None
    if raw_hours not in (None, ""):
        try:
            hours = Decimal(str(raw_hours)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None, "Hours must be a number."
        if not (0 < hours <= MAX_HOURS_PER_ROW):
            return None, f"Hours must be between 0 and {MAX_HOURS_PER_ROW}."
    elif require_hours:
        return None, "Hours are required to approve an entry."

    fields = {
        "work_date": work_date,
        "hours": hours,
        "category": category,
        "activity": activity,
        "project": (data.get("project") or "").strip(),
        "evidence": (data.get("evidence") or "").strip(),
        "degree_link": (data.get("degree_link") or "").strip(),
        "notes": (data.get("notes") or "").strip(),
    }
    limits = {"activity": MAX_ACTIVITY, "project": 120, "evidence": 500, "degree_link": 500, "notes": 500}
    if any(len(fields[k]) > n for k, n in limits.items()):
        return None, "One of the fields is too long."
    return fields, None


def create(fields, status="approved", source="manual"):
    ensure_table()
    return execute(
        "INSERT INTO work_sessions (work_date, hours, category, activity, project, "
        "evidence, degree_link, notes, status, source) "
        "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
        (fields["work_date"], fields["hours"], fields["category"], fields["activity"],
         fields["project"], fields["evidence"], fields["degree_link"], fields["notes"],
         status, source),
    )


def _merge_text(existing, new, sep):
    """Add `new` to `existing` unless it's already there (keeps repeats out)."""
    parts = [p.strip() for p in existing.split(sep) if p.strip()]
    for piece in (p.strip() for p in new.split(sep)):
        if piece and piece not in parts:
            parts.append(piece)
    return sep.join(parts)


def add_draft(fields):
    """
    One draft per project per day. A new Claude entry for a date + project
    that already has a draft is merged into it (activity lines, evidence),
    so the log reads as a daily summary and repeats never pile up.
    Returns the id of the draft that holds the entry.
    """
    ensure_table()
    existing = query(
        "SELECT id, activity, evidence, degree_link, notes FROM work_sessions "
        "WHERE status = 'draft' AND work_date = %s AND project = %s ORDER BY id LIMIT 1",
        (fields["work_date"], fields["project"]), one=True,
    )
    if existing is None:
        return create(fields, status="draft", source="claude")
    execute(
        "UPDATE work_sessions SET activity = %s, evidence = %s, degree_link = %s, notes = %s "
        "WHERE id = %s",
        (_merge_text(existing["activity"], fields["activity"], "\n")[:MAX_ACTIVITY],
         _merge_text(existing["evidence"], fields["evidence"], "; ")[:500],
         _merge_text(existing["degree_link"], fields["degree_link"], "; ")[:500],
         _merge_text(existing["notes"], fields["notes"], "; ")[:500],
         existing["id"]),
    )
    return existing["id"]


def merge_duplicate_drafts():
    """Fold drafts already stored as several rows for the same day + project into one."""
    ensure_table()
    rows = query(
        "SELECT id, work_date, project, activity, evidence, degree_link, notes FROM work_sessions "
        "WHERE status = 'draft' ORDER BY work_date, project, id"
    )
    keep = {}
    for r in rows:
        key = (str(r["work_date"]), r["project"])
        if key not in keep:
            keep[key] = dict(r)
            continue
        k = keep[key]
        k["activity"] = _merge_text(k["activity"], r["activity"], "\n")[:MAX_ACTIVITY]
        for col in ("evidence", "degree_link", "notes"):
            k[col] = _merge_text(k[col], r[col], "; ")[:500]
        k["merged"] = True
        execute("DELETE FROM work_sessions WHERE id = %s", (r["id"],))
    for k in keep.values():
        if k.get("merged"):
            execute(
                "UPDATE work_sessions SET activity = %s, evidence = %s, degree_link = %s, notes = %s "
                "WHERE id = %s",
                (k["activity"], k["evidence"], k["degree_link"], k["notes"], k["id"]),
            )


def update(session_id, fields):
    """Saving from the admin always approves (hours are required there)."""
    ensure_table()
    execute(
        "UPDATE work_sessions SET work_date = %s, hours = %s, category = %s, activity = %s, "
        "project = %s, evidence = %s, degree_link = %s, notes = %s, status = 'approved' "
        "WHERE id = %s",
        (fields["work_date"], fields["hours"], fields["category"], fields["activity"],
         fields["project"], fields["evidence"], fields["degree_link"], fields["notes"],
         session_id),
    )


def delete(session_id):
    ensure_table()
    execute("DELETE FROM work_sessions WHERE id = %s", (session_id,))


def weekly_summary(sessions=None, today=None):
    """Approved hours per week from FIRST_WEEK through the current week, newest first."""
    sessions = list_sessions() if sessions is None else sessions
    today = today or datetime.date.today()
    totals, drafts = {}, {}
    for s in sessions:
        key = s["week_start"]
        if s["status"] == "approved":
            totals[key] = totals.get(key, 0) + (s["hours"] or 0)
        else:
            drafts[key] = drafts.get(key, 0) + 1
    weeks, day = [], FIRST_WEEK
    while day <= week_start(today):
        key = day.isoformat()
        hours = round(totals.get(key, 0), 2)
        weeks.append({
            "week_start": key,
            "hours": hours,
            "drafts": drafts.get(key, 0),
            "meets_minimum": hours >= WEEKLY_MINIMUM,
        })
        day += datetime.timedelta(weeks=1)
    return list(reversed(weeks))


def export_xlsx():
    """Excel copy for the OPT evidence folder: Work Log + Weekly Summary tabs."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    sessions = list_sessions()
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2E8B8B")  # --color-accent-dark (print surface)

    log = wb.active
    log.title = "Work Log"
    log.append(["Date", "Week Start (Mon)", "Hours", "Category", "Activity Description",
                "Project", "Evidence / Reference", "Degree link (MBA-MIS)", "Notes", "Status", "Source"])
    for s in sorted(sessions, key=lambda x: (x["work_date"], x["id"])):
        log.append([_as_date(s["work_date"]), _as_date(s["week_start"]), s["hours"], s["category"],
                    s["activity"], s["project"], s["evidence"], s["degree_link"], s["notes"],
                    s["status"], s["source"]])
    for width, col in zip((12, 16, 8, 28, 60, 22, 32, 40, 30, 10, 10), "ABCDEFGHIJK"):
        log.column_dimensions[col].width = width
    for row in log.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    summary = wb.create_sheet("Weekly Summary")
    summary.append(["Week Starting (Mon)", "Approved Hours", f"Meets {WEEKLY_MINIMUM}-hr Minimum?",
                    "Drafts awaiting approval"])
    for w in reversed(weekly_summary(sessions)):
        summary.append([_as_date(w["week_start"]), w["hours"],
                        "" if w["hours"] == 0 else ("Yes" if w["meets_minimum"] else "NO — review"),
                        w["drafts"] or ""])
    for width, col in zip((20, 16, 22, 24), "ABCD"):
        summary.column_dimensions[col].width = width
    for row in summary.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    for ws in (log, summary):
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
